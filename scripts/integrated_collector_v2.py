# scripts/integrated_collector_v2.py
import json
import logging
import os
import re
from datetime import datetime
from typing import Dict, Optional, Any

import pandas as pd

try:
    import duckdb
except Exception:
    duckdb = None

try:
    import sys as _sys
    import os as _os
    # Garante que o root do projeto está no path para importar db_neon
    _root = _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
    if _root not in _sys.path:
        _sys.path.insert(0, _root)
    import db_neon
    _NEON_OK = db_neon.is_configured()
except Exception:
    db_neon = None  # type: ignore
    _NEON_OK = False

logger = logging.getLogger(__name__)


class KintuadiIntegratedCollectorV2:
    """
    Integrador central de dados – Kintuadi Energy v2

    Responsabilidades:
    - Orquestrar coleta ONS + CCEE
    - Persistir dados em DuckDB para consultas analíticas de baixo uso de memória
    - Salvar snapshot leve para o dashboard/pipeline
    - NÃO realizar análise de mercado
    """

    def __init__(self):
        self.db_path = "data/kintuadi.duckdb"

        # Garantir diretórios antes de registrar FileHandler
        os.makedirs("data", exist_ok=True)
        os.makedirs("logs", exist_ok=True)

        logging.basicConfig(
            level=logging.INFO,
            format="%(asctime)s - %(levelname)s - %(message)s",
            handlers=[
                logging.StreamHandler(),
                logging.FileHandler("logs/kintuadi.log"),
            ],
        )

        try:
            from .ons_collector_v2 import ONSCollectorV2
            from .ccee_collector_v2 import CCEEPLDCollector

            self.ons_collector = ONSCollectorV2(
                username=os.getenv("ONS_API_USER"),
                password=os.getenv("ONS_API_PASSWORD"),
                enable_audit=True,
            )

            self.ccee_collector = CCEEPLDCollector(
                cache_ttl_minutes=60, enable_audit=True
            )

            self.modules_loaded = True
        except Exception as e:
            logger.error(f"Erro ao carregar coletores: {e}")
            self.modules_loaded = False

    def _sanitize_table_name(self, dataset_name: str) -> str:
        table_name = re.sub(r"_(\d{4})(-\d{2})?$", "", dataset_name)
        table_name = table_name.lower()
        table_name = re.sub(r"[^a-z0-9_]", "", table_name)
        return table_name

    def _extract_year_month(self, dataset_name: str) -> tuple[int, Optional[int]]:
        year = 0
        month = None
        match_month = re.search(r"(\d{4})-(\d{2})$", dataset_name)
        if match_month:
            return int(match_month.group(1)), int(match_month.group(2))
        match_year = re.search(r"(\d{4})$", dataset_name)
        if match_year:
            year = int(match_year.group(1))
        return year, month

    def _read_source_file(self, file_path: str) -> pd.DataFrame:
        """
        Lê arquivos ONS com suporte a:
        - XLSX de aba única
        - XLSX com múltiplas abas (ex.: Geracao_Usina_Horaria_2021.xlsx)
        - CSV com detecção simples de encoding
        """
        path_lower = str(file_path).lower()
        try:
            if path_lower.endswith(".xlsx"):
                import openpyxl

                wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                sheet_names = list(wb.sheetnames)
                wb.close()

                if len(sheet_names) <= 1:
                    return pd.read_excel(file_path, engine="openpyxl", sheet_name=0)

                logger.info(
                    "XLSX com múltiplas abas detectado: %s abas em %s",
                    len(sheet_names),
                    os.path.basename(file_path),
                )
                frames = []
                for sheet_name in sheet_names:
                    try:
                        df_sheet = pd.read_excel(
                            file_path,
                            engine="openpyxl",
                            sheet_name=sheet_name,
                        )
                        if not df_sheet.empty:
                            frames.append(df_sheet)
                            logger.info(
                                "  Aba '%s': %s linhas",
                                sheet_name,
                                f"{len(df_sheet):,}",
                            )
                    except Exception as exc:
                        logger.warning(
                            "Falha ao ler aba '%s' de %s: %s",
                            sheet_name,
                            os.path.basename(file_path),
                            exc,
                        )
                if not frames:
                    return pd.DataFrame()
                return pd.concat(frames, ignore_index=True)

            for encoding in ("utf-8-sig", "latin-1", "cp1252"):
                try:
                    return pd.read_csv(
                        file_path,
                        sep=None,
                        engine="python",
                        on_bad_lines="skip",
                        encoding=encoding,
                    )
                except UnicodeDecodeError:
                    continue
        except Exception as exc:
            logger.error("Erro ao ler arquivo fonte %s: %s", file_path, exc)
        return pd.DataFrame()

    def _should_persist_dataset(self, dataset_name: str, persist_mode: str) -> bool:
        """
        Regras de persistência:
        - full: persiste tudo
        - incremental: persiste apenas datasets que podem mudar no tempo:
            * anualizado do ano corrente (YYYY)
            * mensal do mês corrente (YYYY-MM)
        """
        mode = (persist_mode or "full").strip().lower()
        if mode == "full":
            return True

        year, month = self._extract_year_month(dataset_name)
        now = datetime.now()
        current_year = now.year
        current_month = now.month

        if year != current_year:
            return False

        if month is None:
            return True

        return int(month) == int(current_month)

    # ── Neon PostgreSQL ───────────────────────────────────────────────────────

    def _persist_ons_neon(self, dataset_name: str, file_path: str) -> None:
        """Persiste dataset ONS no Neon PostgreSQL."""
        if not _NEON_OK:
            return
        if not os.path.exists(file_path):
            return

        # Apenas tabelas que o app usa
        SUPPORTED = {
            "geracao_usina_horaria", "curva_carga", "despacho_gfom",
            "cmo", "capacidade_instalada", "disponibilidade_usina",
            "ear_diario_subsistema", "ena_diario_subsistema",
            "cvu_usina_termica", "intercambio",
        }
        table_name = self._sanitize_table_name(dataset_name)
        if table_name not in SUPPORTED:
            return

        year, month = self._extract_year_month(dataset_name)
        try:
            is_xlsx = str(file_path).lower().endswith(".xlsx")
            df_src = self._read_source_file(file_path)
            if df_src.empty:
                return

            df_src.columns = [str(c).strip().lower() for c in df_src.columns]
            df_src["ano"] = year
            df_src["mes"] = month

            if not db_neon.table_exists(table_name):
                logger.warning(f"Neon: tabela {table_name} nao existe. Rode: python db_neon.py --create-tables")
                return

            # Deletar periodo antes de reinserir (idempotente)
            if month is not None:
                db_neon.execute(f'DELETE FROM "{table_name}" WHERE ano = %s AND mes = %s', [year, month])
            elif year is not None:
                db_neon.execute(f'DELETE FROM "{table_name}" WHERE ano = %s', [year])

            # Filtrar colunas que existem no destino
            dest_df = db_neon.fetchdf(
                "SELECT column_name FROM information_schema.columns "
                "WHERE table_schema='public' AND table_name=%s", [table_name]
            )
            if dest_df.empty:
                return
            dest_cols = set(dest_df["column_name"].str.lower())
            cols = [c for c in df_src.columns if c in dest_cols]
            if not cols:
                return

            df_insert = df_src[cols].where(pd.notnull(df_src[cols]), None)

            import psycopg2.extras
            col_names = ", ".join([f'"{c}"' for c in cols])
            placeholders = ", ".join(["%s"] * len(cols))
            sql = (f'INSERT INTO "{table_name}" ({col_names}) VALUES ({placeholders}) '
                   f'ON CONFLICT DO NOTHING')
            data = [tuple(r) for r in df_insert.itertuples(index=False, name=None)]
            with db_neon.get_conn() as conn:
                psycopg2.extras.execute_batch(conn.cursor(), sql, data, page_size=500)
            logger.info(f"Neon: {table_name} <- {len(df_insert)} linhas (ano={year}, mes={month})")
        except Exception as e:
            logger.error(f"Neon persist error {dataset_name}: {e}")

    def _persist_pld_neon(self, df: pd.DataFrame, year: int) -> None:
        """Persiste PLD historico no Neon."""
        if not _NEON_OK:
            return
        try:
            dfx = df.copy()
            dfx.columns = [str(c).lower() for c in dfx.columns]

            if "data" not in dfx.columns:
                if {"mes_referencia", "dia", "hora"}.issubset(dfx.columns):
                    mr = pd.to_numeric(dfx["mes_referencia"], errors="coerce").astype("Int64").astype(str).str.zfill(6)
                    dfx["data"] = pd.to_datetime(
                        mr.str[:4] + "-" + mr.str[4:6] + "-" +
                        pd.to_numeric(dfx["dia"], errors="coerce").fillna(1).astype(int).astype(str).str.zfill(2) +
                        " " + pd.to_numeric(dfx["hora"], errors="coerce").fillna(0).astype(int).astype(str).str.zfill(2) + ":00:00",
                        errors="coerce"
                    )
            else:
                dfx["data"] = pd.to_datetime(dfx["data"], errors="coerce")

            if "pld" not in dfx.columns and "pld_hora" in dfx.columns:
                dfx["pld"] = pd.to_numeric(dfx["pld_hora"], errors="coerce")
            if "submercado" in dfx.columns:
                dfx["submercado"] = dfx["submercado"].astype(str).str.upper().str.strip()

            dfx = dfx.dropna(subset=["data", "submercado", "pld"])
            if dfx.empty:
                return

            dfx["ano"]  = dfx["data"].dt.year
            dfx["mes"]  = dfx["data"].dt.month
            dfx["hora"] = dfx["data"].dt.hour
            if "dia" not in dfx.columns:
                dfx["dia"] = dfx["data"].dt.day
            if "mes_referencia" not in dfx.columns:
                dfx["mes_referencia"] = dfx["data"].dt.strftime("%Y%m").astype(int)
            if "periodo_comercializacao" not in dfx.columns:
                dfx["periodo_comercializacao"] = None

            db_neon.execute("DELETE FROM pld_historical WHERE ano = %s", [year])

            keep = ["data","submercado","pld","ano","mes","hora","dia","mes_referencia","periodo_comercializacao"]
            dfx = dfx[[c for c in keep if c in dfx.columns]].where(pd.notnull(dfx), None)

            import psycopg2.extras
            cols = list(dfx.columns)
            col_names = ", ".join([f'"{c}"' for c in cols])
            placeholders = ", ".join(["%s"] * len(cols))
            sql = (f'INSERT INTO pld_historical ({col_names}) VALUES ({placeholders}) '
                   f'ON CONFLICT (data, submercado) DO UPDATE SET pld = EXCLUDED.pld')
            data = [tuple(r) for r in dfx.itertuples(index=False, name=None)]
            with db_neon.get_conn() as conn:
                psycopg2.extras.execute_batch(conn.cursor(), sql, data, page_size=500)
            logger.info(f"Neon: pld_historical <- {len(dfx)} linhas (ano={year})")
        except Exception as e:
            logger.error(f"Neon PLD persist error: {e}")

    def _persist_ons_dataset(self, dataset_name: str, file_path: str):
        if duckdb is None:
            return
        if not os.path.exists(file_path):
            return

        con = duckdb.connect(self.db_path)
        try:
            year, month = self._extract_year_month(dataset_name)
            table_name = self._sanitize_table_name(dataset_name)
            logger.info(f"Persistindo {dataset_name} -> {table_name} (ano={year}, mes={month})")

            is_xlsx = str(file_path).lower().endswith(".xlsx")

            if is_xlsx:
                # XLSX (ex.: Curva_Carga): usar pandas para inferência e insert robusto
                df_src = self._read_source_file(file_path)
                if df_src.empty:
                    return
                con.register("df_src_tmp", df_src)

                # === EVOLUÇÃO AUTOMÁTICA DE SCHEMA (XLSX) ===
                src_info = con.execute("PRAGMA table_info('df_src_tmp')").fetchall()
                src_cols_map = {str(c[1]).strip().lower(): str(c[1]) for c in src_info}

                con.execute(
                    f"""
                    CREATE TABLE IF NOT EXISTS {table_name} AS
                    SELECT * FROM df_src_tmp LIMIT 0
                    """
                )

                existing_cols_info = con.execute(f"PRAGMA table_info('{table_name}')").fetchall()
                existing_cols = {str(c[1]).strip().lower() for c in existing_cols_info}

                for src_col_lower, src_col_original in src_cols_map.items():
                    if src_col_lower not in existing_cols:
                        logger.info(f"Adicionando nova coluna '{src_col_original}' em {table_name}")
                        con.execute(f'ALTER TABLE {table_name} ADD COLUMN "{src_col_original}" VARCHAR')
                # ============================================
            else:
                con.execute(
                    f"""
                    CREATE TABLE IF NOT EXISTS {table_name} AS
                    SELECT *
                    FROM read_csv_auto(?, sample_size=-1, ignore_errors=true)
                    LIMIT 0
                    """,
                    [file_path],
                )

            for stmt in [
                f"ALTER TABLE {table_name} ADD COLUMN ano INTEGER",
                f"ALTER TABLE {table_name} ADD COLUMN mes INTEGER",
            ]:
                try:
                    con.execute(stmt)
                except Exception:
                    pass

            if month is not None:
                # Para séries que migraram de anual para mensal (ex.: GFOM),
                # remover também restos anuais (mes IS NULL) do mesmo ano.
                if table_name in {"despacho_gfom"}:
                    con.execute(
                        f"DELETE FROM {table_name} WHERE ano=? AND (mes=? OR mes IS NULL)",
                        [year, month],
                    )
                else:
                    con.execute(f"DELETE FROM {table_name} WHERE ano=? AND mes=?", [year, month])
            else:
                con.execute(f"DELETE FROM {table_name} WHERE ano=?", [year])

            cols = con.execute(f"PRAGMA table_info('{table_name}')").fetchall()
            col_names = [c[1] for c in cols if c[1] not in ("ano", "mes")]
            col_list = ", ".join([f'"{c}"' for c in col_names])

            if is_xlsx:
                # XLSX com schema variável: alinhar por nome de coluna para evitar BinderError.
                src_info = con.execute("PRAGMA table_info('df_src_tmp')").fetchall()
                src_cols_map = {str(c[1]).strip().lower(): str(c[1]) for c in src_info}

                select_expr = []
                for col in col_names:
                    src_col = src_cols_map.get(str(col).strip().lower())
                    if src_col:
                        select_expr.append(f'"{src_col}"')
                    else:
                        select_expr.append(f'NULL AS "{col}"')
                select_cols = ", ".join(select_expr)

                if month is not None:
                    con.execute(
                        f"""
                        INSERT INTO {table_name}
                        ({col_list}, ano, mes)
                        SELECT {select_cols}, ?, ?
                        FROM df_src_tmp
                        """,
                        [year, month],
                    )
                else:
                    con.execute(
                        f"""
                        INSERT INTO {table_name}
                        ({col_list}, ano, mes)
                        SELECT {select_cols}, ?, NULL
                        FROM df_src_tmp
                        """,
                        [year],
                    )
            else:
                # CSV com schema variável (ex.: GFOM anual vs mensal):
                # criar staging e alinhar colunas por nome para evitar BinderError.
                con.execute(
                    """
                    CREATE OR REPLACE TEMP TABLE _src_csv_tmp AS
                    SELECT * FROM read_csv_auto(?, sample_size=-1, ignore_errors=true)
                    """,
                    [file_path],
                )

                # === EVOLUÇÃO AUTOMÁTICA DE SCHEMA (CSV) ===
                src_info = con.execute("PRAGMA table_info('_src_csv_tmp')").fetchall()
                src_cols_map = {str(c[1]).strip().lower(): str(c[1]) for c in src_info}

                existing_cols_info = con.execute(f"PRAGMA table_info('{table_name}')").fetchall()
                existing_cols = {str(c[1]).strip().lower() for c in existing_cols_info}

                for src_col_lower, src_col_original in src_cols_map.items():
                    if src_col_lower not in existing_cols:
                        logger.info(f"Adicionando nova coluna '{src_col_original}' em {table_name}")
                        con.execute(f'ALTER TABLE {table_name} ADD COLUMN "{src_col_original}" VARCHAR')
                # ===========================================

                select_expr = []
                for col in col_names:
                    src_col = src_cols_map.get(str(col).strip().lower())
                    if src_col:
                        select_expr.append(f'"{src_col}"')
                    else:
                        select_expr.append(f'NULL AS "{col}"')

                select_cols = ", ".join(select_expr)

                if month is not None:
                    con.execute(
                        f"""
                        INSERT INTO {table_name}
                        ({col_list}, ano, mes)
                        SELECT {select_cols}, ?, ?
                        FROM _src_csv_tmp
                        """,
                        [year, month],
                    )
                else:
                    con.execute(
                        f"""
                        INSERT INTO {table_name}
                        ({col_list}, ano, mes)
                        SELECT {select_cols}, ?, NULL
                        FROM _src_csv_tmp
                        """,
                        [year],
                    )

        except Exception as e:
            logger.error(f"Erro ao persistir {dataset_name} no DuckDB: {e}")
        finally:
            con.close()

    def _persist_pld_duckdb(self, df: pd.DataFrame, year: int):
        if duckdb is None:
            return

        con = duckdb.connect(self.db_path)
        try:
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS pld_historical (
                    data TIMESTAMP,
                    submercado VARCHAR,
                    pld DOUBLE,
                    ano INTEGER,
                    mes INTEGER,
                    hora INTEGER,
                    dia INTEGER,
                    mes_referencia INTEGER,
                    periodo_comercializacao INTEGER
                )
                """
            )

            for stmt in [
                "ALTER TABLE pld_historical ADD COLUMN dia INTEGER",
                "ALTER TABLE pld_historical ADD COLUMN mes_referencia INTEGER",
                "ALTER TABLE pld_historical ADD COLUMN periodo_comercializacao INTEGER",
            ]:
                try:
                    con.execute(stmt)
                except Exception:
                    pass

            dfx = df.copy()
            dfx.columns = [str(c).lower() for c in dfx.columns]

            # Normalização de layout CCEE (DIA,HORA,MES_REFERENCIA,PLD_HORA,SUBMERCADO)
            if "data" not in dfx.columns:
                if {"mes_referencia", "dia", "hora"}.issubset(dfx.columns):
                    mr = pd.to_numeric(dfx["mes_referencia"], errors="coerce").astype("Int64").astype(str).str.zfill(6)
                    dfx["data"] = pd.to_datetime(
                        mr.str[:4]
                        + "-"
                        + mr.str[4:6]
                        + "-"
                        + pd.to_numeric(dfx["dia"], errors="coerce").fillna(1).astype(int).astype(str).str.zfill(2)
                        + " "
                        + pd.to_numeric(dfx["hora"], errors="coerce").fillna(0).astype(int).astype(str).str.zfill(2)
                        + ":00:00",
                        errors="coerce",
                    )
            else:
                dfx["data"] = pd.to_datetime(dfx["data"], errors="coerce")

            if "pld" not in dfx.columns and "pld_hora" in dfx.columns:
                dfx["pld"] = pd.to_numeric(dfx["pld_hora"], errors="coerce")
            else:
                dfx["pld"] = pd.to_numeric(dfx.get("pld"), errors="coerce")

            if "submercado" in dfx.columns:
                dfx["submercado"] = dfx["submercado"].astype(str).str.upper().str.strip()
            else:
                dfx["submercado"] = None

            # Campos de granularidade original da CCEE
            if "dia" in dfx.columns:
                dfx["dia"] = pd.to_numeric(dfx["dia"], errors="coerce")
            else:
                dfx["dia"] = pd.to_datetime(dfx["data"], errors="coerce").dt.day

            if "mes_referencia" in dfx.columns:
                dfx["mes_referencia"] = pd.to_numeric(dfx["mes_referencia"], errors="coerce")
            else:
                dfx["mes_referencia"] = pd.to_datetime(dfx["data"], errors="coerce").dt.strftime("%Y%m").astype(float)

            if "periodo_comercializacao" in dfx.columns:
                dfx["periodo_comercializacao"] = pd.to_numeric(dfx["periodo_comercializacao"], errors="coerce")
            else:
                dfx["periodo_comercializacao"] = None

            dfx = dfx.dropna(subset=["data", "submercado", "pld"])
            if dfx.empty:
                return

            dfx["ano"] = dfx["data"].dt.year
            dfx["mes"] = dfx["data"].dt.month
            dfx["hora"] = dfx["data"].dt.hour

            con.execute("DELETE FROM pld_historical WHERE ano = ?", [year])
            con.register("df_temp", dfx)
            con.execute(
                """
                INSERT INTO pld_historical
                SELECT data, submercado, CAST(pld AS DOUBLE), ano, mes, hora,
                       CAST(dia AS INTEGER), CAST(mes_referencia AS INTEGER), CAST(periodo_comercializacao AS INTEGER)
                FROM df_temp
                """
            )
        except Exception as e:
            logger.error(f"Erro ao persistir PLD {year} no DuckDB: {e}")
        finally:
            con.close()

    def _consolidate_ons_audit(self, ons_data: Dict):
        base_path = os.path.join("data", "audit")
        os.makedirs(base_path, exist_ok=True)

        datasets = ons_data.get("datasets", [])
        annual_data = {}

        for ds in datasets:
            name = ds.get("dataset", "")
            file = ds.get("file")
            if not file or not os.path.exists(file):
                continue

            year = None
            if "-" in name:
                try:
                    year = name.split("_")[-1].split("-")[0]
                except Exception:
                    pass
            if year is None:
                parts = name.split("_")
                if parts and parts[-1].isdigit() and len(parts[-1]) == 4:
                    year = parts[-1]
            if not year:
                continue

            annual_data.setdefault(year, {})
            annual_data[year][name] = file

        for year, content in annual_data.items():
            path = os.path.join(base_path, f"ons_{year}.json")
            with open(path, "w", encoding="utf-8") as f:
                json.dump(content, f, indent=2, ensure_ascii=False)

    def collect_all(self, persist_mode: str = "full") -> Optional[Dict[str, Any]]:
        if not self.modules_loaded:
            logger.error("Coletores não carregados.")
            return None

        logger.info("=" * 70)
        logger.info("⚡ KINTUADI ENERGY – DATA COLLECTION v2.0")
        logger.info("=" * 70)

        start_time = datetime.now()

        try:
            results = {
                "metadata": {
                    "collection_start": start_time.isoformat(),
                    "version": "2.0",
                    "project": "Kintuadi Energy Intelligence",
                    "duckdb_enabled": duckdb is not None,
                    "duckdb_path": self.db_path,
                    "persist_mode": (persist_mode or "full").strip().lower(),
                },
                "sources": {},
            }

            logger.info("[1/2] Coletando dados do ONS...")
            ons_data = self.ons_collector.collect_open_data()
            results["sources"]["ons"] = self._normalize_source(ons_data)

            if duckdb is not None:
                for ds in (ons_data or {}).get("datasets", []):
                    dataset_name = ds.get("dataset")
                    file_path = ds.get("file")
                    if dataset_name and file_path and os.path.exists(file_path):
                        if self._should_persist_dataset(dataset_name, persist_mode):
                            self._persist_ons_dataset(dataset_name, file_path)
                            self._persist_ons_neon(dataset_name, file_path)
                        else:
                            logger.info(f"Pulando persist ONS (modo incremental): {dataset_name}")

            logger.info("[2/2] Coletando dados da CCEE...")
            ccee_data = {}

            # Compat: coletores antigos podem não ter collect_pld_historical
            if hasattr(self.ccee_collector, "collect_pld_historical"):
                pld_hist = self.ccee_collector.collect_pld_historical()
            else:
                logger.warning("CCEE collector sem collect_pld_historical; usando collect_pld_data como fallback.")
                pld_hist = {"datasets": []}
                pld_data = self.ccee_collector.collect_pld_data(days=90)
                records = pld_data.get("data", []) if isinstance(pld_data, dict) else []
                if records:
                    df = pd.DataFrame(records)
                    if "mes_referencia" in df.columns:
                        df["year"] = pd.to_numeric(df["mes_referencia"], errors="coerce").astype("Int64").astype(str).str[:4]
                    elif "MES_REFERENCIA" in df.columns:
                        df["year"] = pd.to_numeric(df["MES_REFERENCIA"], errors="coerce").astype("Int64").astype(str).str[:4]
                    else:
                        df["year"] = datetime.now().strftime("%Y")
                    for y, g in df.groupby("year"):
                        pld_hist["datasets"].append({"year": int(y), "records": g.to_dict(orient="records")})

            if duckdb is not None:
                for ds in pld_hist.get("datasets", []):
                    year = ds.get("year")
                    file_path = ds.get("file")
                    try:
                        if (persist_mode or "full").strip().lower() == "incremental":
                            if int(year) != datetime.now().year:
                                logger.info(f"Pulando persist PLD (modo incremental): {year}")
                                continue
                        if file_path and os.path.exists(file_path):
                            df_year = pd.read_csv(file_path)
                        else:
                            recs = ds.get("records") or ds.get("data") or ds.get("timeseries") or []
                            df_year = pd.DataFrame(recs)
                        if df_year.empty:
                            continue
                        logger.info(f"Persistindo PLD {year} no DuckDB + Neon...")
                        yr = int(year) if year else datetime.now().year
                        self._persist_pld_duckdb(df_year, yr)
                        self._persist_pld_neon(df_year, yr)
                    except Exception as e:
                        logger.warning(f"Falha ao persistir PLD {year}: {e}")

            ccee_data["pld_historical"] = self._normalize_source(pld_hist)

            open_data = self.ccee_collector.collect_open_data_csv(limit=100000)
            ccee_data["open_data"] = self._normalize_source(open_data)

            results["sources"]["ccee"] = ccee_data

            end_time = datetime.now()
            results["metadata"]["collection_end"] = end_time.isoformat()
            results["metadata"]["collection_duration"] = (end_time - start_time).total_seconds()
            results["metadata"]["overall_status"] = self._compute_overall_status(results["sources"])

            self._persist(results)
            self._log_summary(results)

            logger.info("✅ Coleta concluída com sucesso.")
            return results

        except Exception as e:
            logger.error(f"Erro crítico na coleta: {e}", exc_info=True)
            return None

    def _normalize_source(self, source_data: Dict) -> Dict:
        normalized = dict(source_data)
        metadata = normalized.get("metadata", {})
        if hasattr(metadata, "to_dict"):
            metadata = metadata.to_dict()
        normalized["metadata"] = metadata
        return normalized

    def _compute_overall_status(self, sources: Dict[str, Dict]) -> str:
        success = 0
        for src in sources.values():
            meta = src.get("metadata", {})
            if meta.get("status") == "success":
                success += 1
        if success == len(sources):
            return "success"
        if success > 0:
            return "partial"
        return "error"

    def _persist(self, data: Dict):
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        complete_file = f"data/kintuadi_raw_{timestamp}.json"
        latest_file = "data/kintuadi_latest.json"

        with open(complete_file, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False, default=str)
        with open(latest_file, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False, default=str)

        logger.info("📁 Dados salvos:")
        logger.info(f"  • {complete_file}")
        logger.info(f"  • {latest_file}")

    def _log_summary(self, data: Dict):
        try:
            ons_meta = data.get("sources", {}).get("ons", {}).get("metadata", {})
            logger.info(
                f"ONS | Datasets coletados: {ons_meta.get('datasets_collected', 'N/A')} | "
                f"Status: {ons_meta.get('status', 'N/A')}"
            )

            ccee_src = data.get("sources", {}).get("ccee", {})
            pld_hist = ccee_src.get("pld_historical", {})
            for ds in pld_hist.get("datasets", []):
                year = ds.get("year")
                stats = ds.get("statistics", {})
                logger.info(
                    f"CCEE | {year} | PLD médio: {stats.get('pld_medio', 'N/A')} | "
                    f"Volatilidade: {stats.get('pld_std', 'N/A')}"
                )
        except Exception as e:
            logger.warning(f"Erro ao gerar resumo: {e}")

    def quick_collect(self, persist_mode: str = "full"):
        print(f"🚀 Iniciando coleta Kintuadi Energy v2... (persist_mode={persist_mode})")
        results = self.collect_all(persist_mode=persist_mode)
        if not results:
            print("❌ Falha na coleta.")
            return None

        print("✅ Coleta concluída.")
        print("📁 Dados disponíveis em data/kintuadi_latest.json")
        print("🌐 Execute: streamlit run dashboard_integrado.py")
        return results
